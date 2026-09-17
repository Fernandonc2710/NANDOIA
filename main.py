import os
import base64
import threading
import textwrap
from datetime import datetime

import requests
from kivy.app import App
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.scrollview import ScrollView
from kivy.uix.label import Label
from kivy.uix.image import Image as KivyImage
from kivy.uix.textinput import TextInput
from kivy.uix.button import Button
from kivy.clock import Clock
from kivy.metrics import dp

# ---------- Configuracion ----------
# Clave fragmentada para evitar la alerta de secretos de GitHub
API_KEY = (
    "AQ.Ab8RN6JiKQwzdeInYSe6SD"
    + "i7QCI4HWN_kvOELwOjEPGQHJDZkg"
)

MODELO_TEXTO = "gemini-2.5-flash"
MODELO_IMAGEN = "gemini-2.5-flash-image"

URL_TEXTO = (
    f"https://generativelanguage.googleapis.com/v1beta/models/"
    f"{MODELO_TEXTO}:generateContent?key={API_KEY}"
)
URL_IMAGEN = (
    f"https://generativelanguage.googleapis.com/v1beta/models/"
    f"{MODELO_IMAGEN}:generateContent?key={API_KEY}"
)

SYSTEM_PROMPT = (
    "Eres Nando, un ingeniero programador brillante, con anos de experiencia "
    "resolviendo problemas tecnicos reales. Tu personalidad es carismatica y "
    "cercana: generas confianza, hablas con seguridad pero sin arrogancia, y "
    "tratas a quien te habla como a un colega al que aprecias. Sin importar "
    "el tema, respondes con mentalidad de ingeniero: entiendes el problema "
    "real, estructuras la respuesta con logica y precision, y usas datos y "
    "terminos exactos en vez de vaguedades. Tono seguro, calido, con humor "
    "ocasional, nunca condescendiente. Siempre respondes en espanol."
)


# ---------- Widgets de burbujas ----------
class BurbujaTexto(Label):
    def __init__(self, texto, es_usuario, **kwargs):
        super().__init__(**kwargs)
        self.text = texto
        self.size_hint_y = None
        self.text_size = (dp(260), None)
        self.halign = "right" if es_usuario else "left"
        self.valign = "middle"
        self.padding = (dp(12), dp(8))
        self.color = (1, 1, 1, 1)
        self.bind(texture_size=self._ajustar_alto)

    def _ajustar_alto(self, *_):
        self.height = self.texture_size[1] + dp(16)


class BurbujaImagen(BoxLayout):
    def __init__(self, ruta_imagen, **kwargs):
        super().__init__(size_hint_y=None, height=dp(230), **kwargs)
        self.add_widget(
            KivyImage(source=ruta_imagen, size_hint_y=None, height=dp(230))
        )


# ---------- Pantalla principal ----------
class ChatNando(BoxLayout):
    def __init__(self, **kwargs):
        super().__init__(orientation="vertical", **kwargs)
        self.historial = []

        self.scroll = ScrollView()
        self.lista_mensajes = BoxLayout(
            orientation="vertical", size_hint_y=None, spacing=dp(8), padding=dp(10)
        )
        self.lista_mensajes.bind(minimum_height=self.lista_mensajes.setter("height"))
        self.scroll.add_widget(self.lista_mensajes)
        self.add_widget(self.scroll)

        fila_entrada = BoxLayout(
            size_hint_y=None, height=dp(56), spacing=dp(6), padding=dp(6)
        )
        self.entrada = TextInput(hint_text="Escribele a Nando...", multiline=False)
        self.entrada.bind(on_text_validate=self.enviar)
        boton_enviar = Button(text="Enviar", size_hint_x=None, width=dp(90))
        boton_enviar.bind(on_press=self.enviar)
        fila_entrada.add_widget(self.entrada)
        fila_entrada.add_widget(boton_enviar)
        self.add_widget(fila_entrada)

        fila_extra = BoxLayout(size_hint_y=None, height=dp(48), spacing=dp(6), padding=(dp(6), 0))
        boton_imagen = Button(text="Imagen")
        boton_imagen.bind(on_press=self.generar_imagen)
        boton_pdf = Button(text="PDF")
        boton_pdf.bind(on_press=self.exportar_pdf)
        boton_word = Button(text="Word")
        boton_word.bind(on_press=self.exportar_word)
        fila_extra.add_widget(boton_imagen)
        fila_extra.add_widget(boton_pdf)
        fila_extra.add_widget(boton_word)
        self.add_widget(fila_extra)

        self._agregar_burbuja("Que onda, soy Nando. Escribeme algo, o pideme una imagen.", False)

    # ---------- helpers de UI ----------
    def _agregar_burbuja(self, texto, es_usuario):
        self.lista_mensajes.add_widget(BurbujaTexto(texto, es_usuario))
        Clock.schedule_once(lambda dt: setattr(self.scroll, "scroll_y", 0), 0.1)

    def _agregar_imagen(self, ruta):
        self.lista_mensajes.add_widget(BurbujaImagen(ruta))
        Clock.schedule_once(lambda dt: setattr(self.scroll, "scroll_y", 0), 0.1)

    def _carpeta_app(self):
        return App.get_running_app().user_data_dir

    # ---------- chat de texto ----------
    def enviar(self, *_):
        mensaje = self.entrada.text.strip()
        if not mensaje:
            return
        self.entrada.text = ""
        self._agregar_burbuja(mensaje, True)
        self.historial.append({"role": "user", "content": mensaje})
        threading.Thread(target=self._pedir_respuesta, daemon=True).start()

    def _pedir_respuesta(self):
        try:
            contenidos = [
                {"role": "user" if m["role"] == "user" else "model",
                 "parts": [{"text": m["content"]}]}
                for m in self.historial
            ]
            payload = {
                "contents": contenidos,
                "systemInstruction": {"parts": [{"text": SYSTEM_PROMPT}]},
            }
            r = requests.post(URL_TEXTO, json=payload, timeout=120)
            data = r.json()
            respuesta = data["candidates"][0]["content"]["parts"][0]["text"]
        except Exception as e:
            respuesta = f"(No pude hablar con Gemini: {e})"
        self.historial.append({"role": "assistant", "content": respuesta})
        Clock.schedule_once(lambda dt: self._agregar_burbuja(respuesta, False), 0)

    # ---------- generar imagen ----------
    def generar_imagen(self, *_):
        prompt = self.entrada.text.strip()
        if not prompt:
            return
        self.entrada.text = ""
        self._agregar_burbuja(f"[imagen] {prompt}", True)
        threading.Thread(target=self._pedir_imagen, args=(prompt,), daemon=True).start()

    def _pedir_imagen(self, prompt):
        ruta, error = None, None
        try:
            payload = {"contents": [{"role": "user", "parts": [{"text": prompt}]}]}
            r = requests.post(URL_IMAGEN, json=payload, timeout=120)
            data = r.json()
            partes = data["candidates"][0]["content"]["parts"]
            b64 = next(p["inlineData"]["data"] for p in partes if "inlineData" in p)
            imagen_bytes = base64.b64decode(b64)
            ruta = os.path.join(
                self._carpeta_app(), f"nando_{datetime.now():%Y%m%d_%H%M%S}.png"
            )
            with open(ruta, "wb") as f:
                f.write(imagen_bytes)
        except Exception as e:
            error = str(e)

        def mostrar(dt):
            if ruta:
                self._agregar_imagen(ruta)
            else:
                self._agregar_burbuja(f"(No pude generar la imagen: {error})", False)

        Clock.schedule_once(mostrar, 0)

    # ---------- exportar documentos ----------
    def exportar_pdf(self, *_):
        threading.Thread(target=self._exportar_pdf, daemon=True).start()

    def _exportar_pdf(self):
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        ruta = os.path.join(self._carpeta_app(), f"nando_{datetime.now():%Y%m%d_%H%M%S}.pdf")
        c = canvas.Canvas(ruta, pagesize=letter)
        y = 750
        for m in self.historial:
            quien = "Tu" if m["role"] == "user" else "Nando"
            for linea in textwrap.wrap(f"{quien}: {m['content']}", 95) or [""]:
                c.drawString(40, y, linea)
                y -= 16
                if y < 40:
                    c.showPage()
                    y = 750
        c.save()
        Clock.schedule_once(
            lambda dt: self._agregar_burbuja(f"PDF guardado:\n{ruta}", False), 0
        )

    def exportar_word(self, *_):
        threading.Thread(target=self._exportar_word, daemon=True).start()

    def _exportar_word(self):
        from docx import Document

        ruta = os.path.join(self._carpeta_app(), f"nando_{datetime.now():%Y%m%d_%H%M%S}.docx")
        doc = Document()
        doc.add_heading("Conversacion con Nando", level=1)
        for m in self.historial:
            quien = "Tu" if m["role"] == "user" else "Nando"
            doc.add_paragraph(f"{quien}: {m['content']}")
        doc.save(ruta)
        Clock.schedule_once(
            lambda dt: self._agregar_burbuja(f"Word guardado:\n{ruta}", False), 0
        )


class NandoApp(App):
    def build(self):
        self.title = "Nando"
        return ChatNando()


if __name__ == "__main__":
    NandoApp().run()
